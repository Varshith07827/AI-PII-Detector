"""Content extraction helpers for supported file types."""
from __future__ import annotations

import io
import csv
from typing import BinaryIO

import pypdf
import docx
import openpyxl

SUPPORTED_TYPES = {"pdf", "docx", "csv", "xlsx", "txt"}


def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = filename.lower().split(".")[-1]
    if ext == "pdf":
        return _extract_pdf(file_bytes)
    if ext == "docx":
        return _extract_docx(file_bytes)
    if ext == "csv":
        return _extract_csv(file_bytes)
    if ext == "xlsx":
        return _extract_xlsx(file_bytes)
    return _extract_txt(file_bytes)


def _extract_pdf(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_csv(data: bytes) -> str:
    decoded = data.decode("utf-8", errors="ignore")
    output_lines = []
    for row in csv.reader(io.StringIO(decoded)):
        output_lines.append(", ".join(row))
    return "\n".join(output_lines)


def _extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines = []
    for sheet in wb:
        for row in sheet.iter_rows(values_only=True):
            row_values = [str(cell) for cell in row if cell is not None]
            if row_values:
                lines.append(", ".join(row_values))
    return "\n".join(lines)


def _extract_txt(data: bytes) -> str:
    return data.decode("utf-8", errors="ignore")
